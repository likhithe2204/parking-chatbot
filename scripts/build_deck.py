from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SHOTS = DOCS / "screenshots"
OUT = DOCS / "parking_chatbot_stage1.pptx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x2B, 0x6C, 0xB0)
GREY = RGBColor(0x45, 0x4B, 0x57)
LIGHT = RGBColor(0xF1, 0xF4, 0xF9)
PALE = RGBColor(0xC9, 0xD6, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

REQUIRED = [
    DOCS / "architecture.png",
    SHOTS / "01_ingestion.png",
    SHOTS / "02_chat_info.png",
    SHOTS / "03_reservation_flow.png",
    SHOTS / "04_guardrail.png",
    SHOTS / "05_evaluation.png",
    SHOTS / "06a_tests_passed.png",
    SHOTS / "06b_ci_pipeline.png",
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_text(slide, text, left, top, width, height, size=18, bold=False,
             color=GREY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    return box


def add_bullets(slide, items, left, top, width, height, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.space_after = Pt(12)
        run = para.add_run()
        run.text = f"\u2022  {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = GREY
        run.font.name = FONT


def add_image_fit(slide, path, left, top, max_w, max_h):
    frame = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(max_w), Inches(max_h)
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = LIGHT
    frame.line.fill.background()

    pad = 0.12
    with Image.open(path) as img:
        w, h = img.size
    scale = min((max_w - 2 * pad) / w, (max_h - 2 * pad) / h)
    width, height = w * scale, h * scale
    x = left + (max_w - width) / 2
    y = top + (max_h - height) / 2
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(width), Inches(height))


def add_stat_card(slide, value, label, left, top, width=2.15, height=1.35):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.fill.background()
    add_text(slide, value, left, top + 0.12, width, 0.65, size=30, bold=True,
             color=BLUE, align=PP_ALIGN.CENTER)
    add_text(slide, label, left, top + 0.78, width, 0.5, size=13, align=PP_ALIGN.CENTER)


def add_table(slide, headers, rows, left, top, col_widths, row_height=0.56, size=13):
    shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(left), Inches(top),
        Inches(sum(col_widths)), Inches(row_height * (len(rows) + 1)),
    )
    table = shape.table
    for index, width in enumerate(col_widths):
        table.columns[index].width = Inches(width)

    def fill(cell, text, bold, color, background):
        cell.fill.solid()
        cell.fill.fore_color.rgb = background
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = cell.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = FONT

    for col, header in enumerate(headers):
        fill(table.cell(0, col), header, True, WHITE, NAVY)
    for row_index, row in enumerate(rows, start=1):
        background = LIGHT if row_index % 2 == 0 else WHITE
        for col, value in enumerate(row):
            fill(table.cell(row_index, col), value, col == 0, GREY, background)
    for row in table.rows:
        row.height = Inches(row_height)


def base_slide(title, notes=""):
    slide = prs.slides.add_slide(BLANK)
    add_text(slide, title, 0.6, 0.35, 12, 0.8, size=30, bold=True, color=NAVY)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.15), Inches(1.2), Inches(0.06)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    add_text(slide, "Parking Reservation Chatbot  |  Stage 1", 0.6, 7.0, 8, 0.3, size=11)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def title_slide():
    slide = prs.slides.add_slide(BLANK)
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = NAVY
    background.line.fill.background()
    add_text(slide, "Parking Space Reservation Chatbot", 0.9, 2.3, 11.5, 1.2,
             size=44, bold=True, color=WHITE)
    add_text(slide, "Stage 1 of 4: RAG chatbot, guardrails and evaluation", 0.9, 3.6, 11.5,
             0.7, size=24, color=PALE)
    add_text(slide, "Edupuganti Likhith", 0.9, 5.4, 11.5, 0.5, size=18, color=WHITE)
    add_text(slide, "github.com/likhithe2204/parking-chatbot", 0.9, 5.9, 11.5, 0.5,
             size=16, color=PALE)
    slide.notes_slide.notes_text_frame.text = (
        "A chatbot for a parking facility. It answers questions from real data, "
        "collects reservations, and keeps private data private."
    )


def coverage_slide():
    slide = base_slide(
        "What the task asked and what was delivered",
        "Go row by row. Every requirement of Stage 1 has a working feature and proof "
        "on a later slide.",
    )
    rows = [
        ("RAG architecture", "Retrieve 3 chunks, answer only from them with GPT (EPAM DIAL)", "Slides 3 and 4"),
        ("Vector database", "Chroma with all-MiniLM-L6-v2 embeddings", "Slide 4"),
        ("Static and dynamic data split", "Static data in Chroma, live data in SQLite", "Slide 4"),
        ("Provide information", "Agent chooses vector search or SQL tools", "Slide 5"),
        ("Collect reservation inputs", "One field at a time, each answer validated", "Slide 6"),
        ("Guardrails with NLP models", "Presidio and spaCy mask private data, block injection", "Slide 7"),
        ("Evaluation", "Recall@K, Precision@K, MRR, accuracy and latency report", "Slide 8"),
        ("Extra artifacts", "README, 68 pytest tests, GitHub Actions CI, this deck", "Slide 9"),
    ]
    add_table(slide, ("Requirement", "What was built", "Proof"), rows, 0.6, 1.5,
              (3.6, 6.9, 1.65))


def text_slide(title, bullets, notes=""):
    slide = base_slide(title, notes)
    add_bullets(slide, bullets, 0.6, 1.6, 12.0, 5.2, size=22)


def image_slide(title, bullets, image, notes=""):
    slide = base_slide(title, notes)
    add_bullets(slide, bullets, 0.6, 1.5, 4.3, 5.3, size=17)
    add_image_fit(slide, image, 5.2, 1.5, 7.55, 5.3)


def wide_image_slide(title, image, notes=""):
    slide = base_slide(title, notes)
    add_image_fit(slide, image, 0.6, 1.45, 12.15, 5.4)


def evaluation_slide():
    slide = base_slide(
        "Evaluation",
        "Recall says the right documents were found. MRR says how high the best one "
        "ranks. Precision is lower only because each document has few chunks.",
    )
    add_stat_card(slide, "1.00", "Recall@3", 0.6, 1.5)
    add_stat_card(slide, "0.95", "MRR", 2.95, 1.5)
    add_stat_card(slide, "16 / 16", "Answer checks passed", 0.6, 3.05)
    add_stat_card(slide, "~6 s", "Full reply (mostly LLM)", 2.95, 3.05)
    add_bullets(
        slide,
        [
            "Precision@3 is 0.44: few chunks per document.",
            "Retrieval takes about 8 ms, SQL 0.2 ms.",
        ],
        0.6, 4.7, 4.5, 2.0, size=15,
    )
    add_image_fit(slide, SHOTS / "05_evaluation.png", 5.2, 1.5, 7.55, 5.3)


def tests_slide():
    slide = base_slide(
        "Tests, CI/CD and documentation",
        "Left: pytest result. Right: the green pipeline run on every push.",
    )
    add_bullets(
        slide,
        [
            "68 pytest tests, at least 2 per core module, no API key needed.",
            "GitHub Actions runs them on every push. README documents setup, usage and structure.",
        ],
        0.6, 1.4, 12.0, 1.2, size=17,
    )
    add_image_fit(slide, SHOTS / "06a_tests_passed.png", 0.6, 2.7, 6.0, 4.2)
    add_image_fit(slide, SHOTS / "06b_ci_pipeline.png", 6.75, 2.7, 6.0, 4.2)


def build() -> None:
    missing = [str(p.relative_to(ROOT)) for p in REQUIRED if not p.exists()]
    if missing:
        raise SystemExit("Missing images:\n  " + "\n  ".join(missing))

    title_slide()
    coverage_slide()

    wide_image_slide(
        "Architecture",
        DOCS / "architecture.png",
        "Follow the flow left to right: guardrail, agent, tools, data stores, reply. "
        "Grey boxes are the later stages.",
    )

    image_slide(
        "RAG with a vector database",
        [
            "Static data (info, zones, location, booking rules) is chunked and embedded.",
            "Chunks live in Chroma. A question retrieves the 3 closest.",
            "Live data (hours, prices, free spaces) lives in SQL, because it changes often.",
        ],
        SHOTS / "01_ingestion.png",
        "Static data goes to the vector database, dynamic data to SQL.",
    )

    image_slide(
        "Providing information",
        [
            "The agent picks a tool for each question.",
            "Prices and availability always come from SQL, never from model memory.",
            "Follow-up questions work. Off-topic questions are politely refused.",
        ],
        SHOTS / "02_chat_info.png",
        "Show location, price, availability, then the follow-up 'And Zone A?'.",
    )

    image_slide(
        "Collecting a reservation",
        [
            "Asks for name, car number, start and end, one field at a time.",
            "Each answer is validated at once, for example a car number like TS09AB1234.",
            "Summary and confirmation come first. The bot never says 'confirmed'; the administrator decides in Stage 2.",
        ],
        SHOTS / "03_reservation_flow.png",
        "Point at the rejected car number ABC, then the summary and the final message.",
    )

    image_slide(
        "Guardrails: data protection",
        [
            "Presidio with a spaCy model masks emails, phones and card numbers.",
            "Prompt injection is blocked before it reaches the model.",
            "The same masking runs on replies, so leaked details are blurred.",
        ],
        SHOTS / "04_guardrail.png",
        "Walk through the five inputs, then the output guard on the leaked reply.",
    )

    evaluation_slide()
    tests_slide()

    text_slide(
        "Outcomes and next stages",
        [
            "Working chatbot that informs users and collects reservations.",
            "Data protection on input and output.",
            "Evaluation report with Recall@K, Precision@K, MRR and latency.",
            "Next: Stage 2 administrator agent, Stage 3 MCP server, Stage 4 LangGraph orchestration.",
            "Repository: github.com/likhithe2204/parking-chatbot",
        ],
        "Close with the three outcomes, the roadmap and the repository link.",
    )

    prs.save(OUT)
    print(f"Saved {OUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()